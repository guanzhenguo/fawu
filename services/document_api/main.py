import re
import zipfile
from difflib import SequenceMatcher
from io import BytesIO

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel, Field

app = FastAPI(
    title="FAWU Document and AI Safety Gateway",
    version="0.1.0",
    description="Stateless document extraction, structured diff, and masking preview service.",
)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class ClausePayload(BaseModel):
    code: str = Field(min_length=1, max_length=64)
    content: str = ""
    locked: bool = False
    core: bool = False


class DiffRequest(BaseModel):
    baseline: list[ClausePayload]
    submitted: list[ClausePayload]


class DiffItem(BaseModel):
    code: str
    change: str
    severity: str
    similarity: float
    detail: str


class DiffResponse(BaseModel):
    risk_level: str
    items: list[DiffItem]


class MaskRequest(BaseModel):
    text: str = Field(max_length=200_000)


class MaskResponse(BaseModel):
    masked_text: str
    entity_counts: dict[str, int]
    warning: str


PATTERNS = {
    "CN_ID": re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
    "MOBILE": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "EMAIL": re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"),
    "BANK_CARD": re.compile(r"(?<!\d)\d{16,19}(?!\d)"),
}


@app.get("/health")
def health():
    return {"status": "ok", "service": "fawu-document-api"}


@app.post("/v1/diff/structured", response_model=DiffResponse)
def structured_diff(payload: DiffRequest):
    baseline = {item.code: item for item in payload.baseline}
    submitted = {item.code: item for item in payload.submitted}
    items: list[DiffItem] = []

    for code, base in baseline.items():
        candidate = submitted.get(code)
        severity = "HIGH" if base.locked or base.core else "MEDIUM"
        if candidate is None:
            items.append(
                DiffItem(
                    code=code,
                    change="REMOVED",
                    severity=severity,
                    similarity=0,
                    detail="模板条款在提交稿中缺失。",
                )
            )
            continue
        base_text = normalize(base.content)
        candidate_text = normalize(candidate.content)
        if base_text != candidate_text:
            items.append(
                DiffItem(
                    code=code,
                    change="MODIFIED",
                    severity=severity,
                    similarity=round(SequenceMatcher(None, base_text, candidate_text).ratio(), 4),
                    detail="提交稿与模板基线内容不一致。",
                )
            )

    for code in sorted(set(submitted) - set(baseline)):
        items.append(
            DiffItem(
                code=code,
                change="ADDED",
                severity="MEDIUM",
                similarity=0,
                detail="提交稿新增了模板之外的条款。",
            )
        )

    risk_level = (
        "HIGH" if any(i.severity == "HIGH" for i in items) else "MEDIUM" if items else "LOW"
    )
    return DiffResponse(risk_level=risk_level, items=items)


@app.post("/v1/docx/extract")
async def extract_docx(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=415, detail="仅支持 .docx 文件")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件不能超过 20MB")
    validate_docx_archive(content)
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法解析 DOCX 文件") from exc
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
    ]
    return {"paragraphs": paragraphs, "tables": tables, "paragraph_count": len(paragraphs)}


@app.post("/v1/mask/preview", response_model=MaskResponse)
def mask_preview(payload: MaskRequest):
    text = payload.text
    counts: dict[str, int] = {}
    for entity_type, pattern in PATTERNS.items():
        index = 0

        def replace(match, current_type=entity_type):
            nonlocal index
            index += 1
            return f"<{current_type}_{index}>"

        text, count = pattern.subn(replace, text)
        counts[entity_type] = count
    return MaskResponse(
        masked_text=text,
        entity_counts=counts,
        warning="仅用于预览；生产环境必须将可逆映射加密存入内网密钥域。",
    )


def validate_docx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            if "word/document.xml" not in names:
                raise HTTPException(status_code=422, detail="文件不是有效 DOCX")
            total = sum(item.file_size for item in archive.infolist())
            if total > MAX_UNCOMPRESSED_BYTES:
                raise HTTPException(status_code=413, detail="解压后文件过大")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise HTTPException(status_code=422, detail="不接受包含宏的文档")
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="文件不是有效 DOCX") from exc


def normalize(value: str) -> str:
    return " ".join(value.split())
